"""Batch pipeline: spreadsheet in, papers (and Zotero entries) out.

Design notes worth keeping:

* Batch runs look like crawlers. A per-paper delay plus a stop-loss on repeated
  ``blocked`` results matters more than raw speed — one ACS block can outlive
  the whole run.
* Zotero's connector can only attach files to items created in the same
  session, so an existing DOI cannot receive a new PDF. We detect it first and
  skip, instead of silently creating a duplicate entry.
"""

from __future__ import annotations

import re
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook

from . import zotero
from .fetch import fetch as run_fetch
from .metadata import DOI_RE
from .models import PdfStatus, Report

Logger = Callable[[str], None]

MAX_ZOTERO_ATTACHMENT_MB = zotero.MAX_ATTACHMENT_MB


def _noop(_: str) -> None:
    return None


@dataclass
class Row:
    source_row: int
    doi: str
    title: str = ""
    status: str = "pending"
    detail: str = ""
    workdir: str = ""
    zotero_state: str = ""


@dataclass
class BatchSummary:
    total: int = 0
    downloaded: int = 0
    skipped: int = 0
    failed: int = 0
    partial: int = 0
    no_doi: int = 0
    zotero_added: int = 0
    zotero_skipped: int = 0
    aborted: str = ""
    report_path: str = ""


def _clean_doi(raw: str) -> str:
    return raw.strip().rstrip(".,;)]}>'\"”“’")


_CJK = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uff00-\uffef]")
_DOI_SHAPE = re.compile(r"10\.\d{4,9}/\S+")


def _is_plausible_doi(value: str) -> bool:
    """说明书里常出现 ``10.1038/文章号`` 这类示例，不能当成真 DOI 去下载。"""
    if _CJK.search(value):
        return False
    return bool(_DOI_SHAPE.fullmatch(value))


def read_manifest(path: Path) -> list[Row]:
    """Pull every DOI out of an .xlsx, keeping row numbers and a title hint.

    只在第一个工作表里记录"没有 DOI 的行"：其余工作表通常是同一批文献的
    精选视图或说明页，把它们也列出来只会让结果表变吵。
    """
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    rows: list[Row] = []
    seen: set[str] = set()

    for sheet_index, sheet in enumerate(workbook.worksheets):
        for index, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            dois: list[str] = []
            title_hint = ""

            for cell in cells:
                if not isinstance(cell, str):
                    continue
                text = cell.strip()
                if not text:
                    continue
                found = [_clean_doi(m) for m in DOI_RE.findall(text)]
                found = [d for d in found if _is_plausible_doi(d)]
                if found:
                    dois.extend(found)
                elif len(text) > len(title_hint) and not text.lower().startswith(("http://", "https://")):
                    # 标题通常是最长的那个文本单元格（前面可能是"研究方向"之类的短标签，后面可能是 DOI 网址）
                    title_hint = text

            if not dois:
                # 首个工作表的第一行通常是表头，不计入无 DOI 行
                if sheet_index == 0 and index > 1:
                    rows.append(Row(source_row=index, doi="", title=title_hint, status="no-doi"))
                continue

            for doi in dois:
                key = doi.lower()
                if key in seen:
                    continue
                seen.add(key)
                rows.append(Row(source_row=index, doi=doi, title=title_hint))

    workbook.close()
    return rows


def read_previous_report(path: Path) -> dict[str, dict[str, str]]:
    """doi(lower) -> 上次记录的字段，供 ``--resume`` 恢复状态。

    按表头名定位列，这样旧版本结果表（没有"标题"列）也能读。
    """
    if not path.exists():
        return {}
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    result: dict[str, dict[str, str]] = {}
    header_map: dict[str, int] = {}

    for index, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index == 1:
            header_map = {str(c or "").strip(): i for i, c in enumerate(cells)}
            continue
        if not cells:
            continue

        def field(name: str) -> str:
            position = header_map.get(name)
            if position is None or position >= len(cells):
                return ""
            return str(cells[position] or "").strip()

        doi = field("DOI").lower()
        if doi:
            result[doi] = {
                "status": field("状态"),
                "detail": field("说明"),
                "workdir": field("保存目录"),
                "title": field("标题"),
            }
    workbook.close()
    return result


def write_report(path: Path, rows: Iterable[Row]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    sheet.append(["源行", "DOI", "标题", "状态", "说明", "保存目录", "Zotero"])
    for row in rows:
        sheet.append(
            [
                row.source_row,
                row.doi,
                row.title,
                row.status,
                row.detail,
                row.workdir,
                row.zotero_state,
            ]
        )
    for column, width in zip("ABCDEFG", (6, 36, 52, 12, 40, 58, 22)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))


def _attachment_files(report: Report) -> list[tuple[str, Path, str, str]]:
    """(title, path, url, mime) for the main PDF plus every SI file on disk."""
    items: list[tuple[str, Path, str, str]] = []
    if report.pdf_path:
        main = Path(report.pdf_path)
        if main.exists():
            items.append(("正文 PDF", main, report.landing_url or "", "application/pdf"))
    for asset in report.si_assets or []:
        path = Path(report.workdir) / asset.filename
        if path.exists():
            mime = asset.content_type or "application/octet-stream"
            items.append((asset.filename, path, asset.url, mime))
    return items


def push_to_zotero(report: Report, *, target: str = "") -> tuple[str, int]:
    """Create the item and hang the downloaded files on it. Returns (state, added)."""
    session_id = "doifetch-" + uuid.uuid4().hex[:12]
    zotero.save_item(report.metadata, session_id, uri=report.landing_url or "")

    if target:
        zotero.move_to_target(session_id, target)

    added = 0
    for title, path, url, mime in _attachment_files(report):
        if path.stat().st_size > MAX_ZOTERO_ATTACHMENT_MB * 1024 * 1024:
            continue
        zotero.save_attachment(
            session_id,
            "doifetch-main",
            path.read_bytes(),
            title=title,
            url=url,
            mime=mime,
        )
        added += 1
    return ("已入库", added)


def run_batch(
    rows: list[Row],
    *,
    out: Path | None = None,
    want_si: bool = True,
    delay: float = 3.0,
    limit: int = 0,
    use_zotero: bool = False,
    zotero_target: str = "",
    resume: dict[str, str] | None = None,
    skip_prefixes: list[str] | None = None,
    auth_mode: str | None = None,
    log: Logger | None = None,
) -> BatchSummary:
    log = log or _noop
    summary = BatchSummary(total=len([r for r in rows if r.doi]))
    resume = resume or {}
    skip_prefixes = [p.strip() for p in (skip_prefixes or []) if p.strip()]

    usable = [row for row in rows if row.doi]
    if skip_prefixes:
        kept: list[Row] = []
        for row in usable:
            if any(row.doi.lower().startswith(prefix.lower()) for prefix in skip_prefixes):
                row.status = "skipped"
                row.detail = "按 --skip-prefix 跳过"
                summary.skipped += 1
            else:
                kept.append(row)
        usable = kept
    if limit:
        usable = usable[:limit]

    if use_zotero:
        if zotero.ping():
            log(f"  Zotero 已连接（{zotero.version() or '版本未知'}）")
        else:
            log("  Zotero 没有响应，本次不写入 Zotero")
            use_zotero = False

    # 出版商级节奏：对同一个站点连发太快一定会被限流。实测 34 篇里第 12 篇左右
    # nature.com 就开始拒绝，而"某个出版商被拒"不该拖停整批任务。
    publisher_interval = {"10.1038": 12.0, "10.1016": 15.0, "10.1126": 12.0}
    cooldown_seconds = 240.0
    next_slot: dict[str, float] = {}
    blocked_hits: dict[str, int] = {}
    cooldown_until: dict[str, float] = {}
    given_up: set[str] = set()
    blocked_run = 0

    def publisher_of(doi: str) -> str:
        return (doi.split("/")[0] or doi).lower()

    for position, row in enumerate(usable, start=1):
        previous = resume.get(row.doi.lower()) if resume else None
        if previous:
            row.title = previous.get("title") or row.title
            row.workdir = previous.get("workdir") or ""
            row.detail = previous.get("detail") or ""
            if previous.get("status") == "downloaded":
                row.status = "downloaded"      # 保留"已完成"，否则下次 --resume 会重下
                summary.skipped += 1
                log(f"[{position}/{len(usable)}] {row.doi}  跳过（上次已完成）")
                continue

        key = publisher_of(row.doi)
        if key in given_up:
            row.status = "failed"
            row.detail = f"{key} 持续拒绝，本批已放弃"
            summary.failed += 1
            log(f"[{position}/{len(usable)}] {row.doi}  跳过（{key} 已被连续拒绝，冷却后仍不行）")
            continue

        now = time.monotonic()
        wait = max(next_slot.get(key, 0.0), cooldown_until.get(key, 0.0)) - now
        if wait > 0:
            if wait > 5:
                log(f"   等待 {wait:.0f}s（{key} 节奏控制）")
            time.sleep(min(wait, 300.0))

        log(f"[{position}/{len(usable)}] {row.doi}")

        try:
            report = run_fetch(
                row.doi,
                out=out,
                want_si=want_si,
                auth_mode=auth_mode,
                log=lambda message: log("   " + message),
            )
        except Exception as exc:
            row.status = "failed"
            row.detail = f"{type(exc).__name__}: {exc}"
            summary.failed += 1
            log(f"   失败：{row.detail}")
            next_slot[key] = time.monotonic() + publisher_interval.get(key, delay)
            continue

        next_slot[key] = time.monotonic() + publisher_interval.get(key, delay)

        row.workdir = report.workdir
        row.detail = report.pdf_source or report.pdf_error
        if report.metadata and report.metadata.title:
            row.title = report.metadata.title

        if report.ok:
            row.status = "downloaded"
            summary.downloaded += 1
        elif report.si_assets:
            # 正文没拿到但补充材料下来了，别报成彻底失败
            row.status = "partial"
            row.detail = f"正文失败（{report.pdf_error or report.pdf_status}），SI {len(report.si_assets)} 个"
            summary.partial += 1
        else:
            row.status = "failed"
            summary.failed += 1

        if report.pdf_status == PdfStatus.BLOCKED.value:
            blocked_run += 1
            blocked_hits[key] = blocked_hits.get(key, 0) + 1
            if blocked_hits[key] >= 2:
                if key in cooldown_until:
                    given_up.add(key)
                    log(f"   {key} 冷却后仍被拒绝，本批放弃它的剩余条目")
                else:
                    cooldown_until[key] = time.monotonic() + cooldown_seconds
                    log(f"   {key} 连续被拒，冷却 {cooldown_seconds:.0f}s 后再试")
                blocked_hits[key] = 0
            if blocked_run >= 4:
                summary.aborted = "连续多篇被拒，看起来整体网络受限，已停止"
                log("   " + summary.aborted)
                break
        else:
            blocked_run = 0
            blocked_hits[key] = 0

        if use_zotero and report.ok:
            existing = zotero.find_by_doi(report.metadata.doi or row.doi)
            if existing:
                row.zotero_state = "跳过（库中已有）"
                summary.zotero_skipped += 1
                log("   Zotero：库中已有，跳过建档")
            else:
                try:
                    state, added = push_to_zotero(report, target=zotero_target)
                    row.zotero_state = f"{state}（{added} 个附件）"
                    summary.zotero_added += 1
                    log(f"   Zotero：{row.zotero_state}")
                except zotero.ZoteroError as exc:
                    row.zotero_state = f"失败：{exc}"
                    log(f"   Zotero 失败：{exc}")

    summary.no_doi = len([r for r in rows if not r.doi])
    return summary
